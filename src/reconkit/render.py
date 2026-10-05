"""Renderers: turn a :class:`Report` into Markdown, HTML, DOCX or JSON."""

from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from pathlib import Path
from string import Template

from . import __version__
from .models import Finding, Report, SEVERITIES

TEMPLATE_PATH = Path(__file__).with_name("templates") / "report.md.tmpl"
HTML_TEMPLATE_PATH = Path(__file__).with_name("templates") / "report.html.tmpl"

DISCLAIMER_SHORT = (
    "**Confidential.** This report is prepared for the named client only and describes the security "
    "posture of the systems in scope at the time of testing. It must not be redistributed without "
    "written consent."
)

DISCLAIMER_FULL = """\
This document contains information that could be used to attack the systems described. It is issued
solely to the client named above and is confidential.

The assessment reflects the state of the in-scope systems during the testing window only. Security
posture changes continuously; findings may no longer be accurate after remediation or after
infrastructure changes. Findings are derived from automated analysis of scan output and require
human validation before they are treated as confirmed vulnerabilities.

Testing was performed within the agreed scope and rules of engagement. No action was taken to
modify, delete or exfiltrate data, and no denial-of-service condition was induced.

The tooling used to generate this document is provided as-is, without warranty of any kind. The
authors accept no liability for use outside an authorised engagement.
"""

METHODOLOGY = """\
Data was collected with standard, non-destructive discovery tooling and then normalised and analysed
by `reconkit`.

**Collection**

* TCP port and service discovery: `nmap -sV -sC`
* HTTP probing, fingerprinting and technology detection: `httpx -jsonl`
* Template-based vulnerability detection: `nuclei -jsonl`

**Analysis**

* Port exposure rules (data-driven, see `rules/default.json`)
* Service fingerprint review against script output from nmap NSE
* Web surface review: management interfaces, non-production hostnames, cleartext transport,
  sensitive paths

**Limitations**

* Version fingerprints are recorded for manual review; reconkit does not claim CVE matches.
* No authenticated testing was performed unless stated in the scope above.
* Business-logic, client-side and social-engineering classes of issue are out of scope for automated
  analysis.
* Automated checks produce both false positives and false negatives. Every finding below should be
  validated manually before it is reported to a system owner.
"""


def md_escape(value: str) -> str:
    """Escape pipes and newlines so values survive inside a Markdown table cell."""
    return value.replace("|", "\\|").replace("\n", " ").strip()


def _finding_block(index: int, finding: Finding) -> str:
    lines = [
        "### {}. {} — {}".format(index, finding.severity.upper(), finding.title),
        "",
        "| | |",
        "|---|---|",
        "| **Check ID** | `{}` |".format(finding.check_id),
        "| **Severity** | {} |".format(finding.severity.capitalize()),
        "| **Affected asset** | `{}` |".format(md_escape(finding.asset)),
        "",
    ]
    if finding.description:
        lines += ["**Description**", "", finding.description.strip(), ""]
    if finding.evidence:
        lines += ["**Evidence**", "", "```", finding.evidence.strip(), "```", ""]
    if finding.remediation:
        lines += ["**Remediation**", "", finding.remediation.strip(), ""]
    if finding.references:
        lines += ["**References**", ""]
        lines += ["- {}".format(ref) for ref in finding.references]
        lines += [""]
    return "\n".join(lines)


def _exec_summary(report: Report) -> str:
    counts = report.counts_by_severity()
    total = len(report.findings)
    hosts_up = len(report.hosts)
    open_ports = sum(len(h.open_ports) for h in report.hosts)

    if total == 0:
        headline = (
            "No findings were raised by the automated checks. This does **not** mean the environment "
            "is free of vulnerabilities; it means nothing matched the rule set. Treat the absence of "
            "findings as a prompt for manual testing, not as a clean bill of health."
        )
    else:
        headline = (
            "{total} finding(s) were raised across {hosts} host(s) and {targets} HTTP endpoint(s). "
            "The highest severity observed is **{top}**. Address critical and high findings first: "
            "they represent reachable paths to credential theft or remote code execution."
        ).format(
            total=total,
            hosts=hosts_up,
            targets=len(report.web_targets),
            top=report.findings[0].severity.upper(),
        )

    table = [
        "| Severity | Count |",
        "|---|---|",
    ]
    for severity in SEVERITIES:
        table.append("| {} | {} |".format(severity.capitalize(), counts.get(severity, 0)))
    table.append("| **Total** | **{}** |".format(total))

    assets = [
        "",
        "**Observed attack surface**",
        "",
        "- Hosts reported up: **{}**".format(hosts_up),
        "- Open ports: **{}**".format(open_ports),
        "- HTTP endpoints probed: **{}**".format(len(report.web_targets)),
    ]

    top = [f for f in report.findings if f.severity in ("critical", "high")][:5]
    if top:
        assets += ["", "**Priority items**", ""]
        assets += [
            "{}. `{}` — {} ({})".format(i, f.asset, f.title, f.severity)
            for i, f in enumerate(top, 1)
        ]

    return "\n".join([headline, "", *table, *assets])


def _scope(report: Report) -> str:
    if not report.scope:
        return (
            "Scope was not supplied on the command line. **Populate this section before the report "
            "leaves the team**: list every in-scope host, network range or application, and record "
            "what was explicitly excluded."
        )
    return "\n".join("- `{}`".format(item) for item in report.scope)


def _findings(report: Report) -> str:
    if not report.findings:
        return "_No findings were raised by the automated check set._"
    blocks = []
    for index, finding in enumerate(report.findings, 1):
        blocks.append(_finding_block(index, finding))
    return "\n---\n\n".join(blocks)


def _host_inventory(report: Report) -> str:
    if not report.hosts:
        return "_No hosts were parsed from the supplied scan data._"
    rows = ["| Host | Hostname | Port | Service | Banner |", "|---|---|---|---|---|"]
    for host in report.hosts:
        if not host.open_ports:
            rows.append(
                "| {} | {} | – | – | no open ports observed |".format(
                    md_escape(host.address), md_escape(host.hostname) or "–"
                )
            )
            continue
        for port in host.open_ports:
            rows.append(
                "| {} | {} | {}/{} | {} | {} |".format(
                    md_escape(host.address),
                    md_escape(host.hostname) or "–",
                    port.number,
                    port.protocol,
                    md_escape(port.service) or "unknown",
                    md_escape(port.banner),
                )
            )
    return "\n".join(rows)


def _web_surface(report: Report) -> str:
    if not report.web_targets:
        return "_No HTTP endpoints were parsed from the supplied httpx data._"
    rows = ["| URL | Status | Title | Server | Technologies |", "|---|---|---|---|---|"]
    for target in report.web_targets:
        rows.append(
            "| {} | {} | {} | {} | {} |".format(
                md_escape(target.url),
                target.status_code if target.status_code is not None else "–",
                md_escape(target.title) or "–",
                md_escape(target.webserver) or "–",
                md_escape(", ".join(target.technologies)) or "–",
            )
        )
    return "\n".join(rows)


def _manual_review(report: Report) -> str:
    if not report.manual_review:
        return "_No version fingerprints were collected, so there is nothing queued for manual review._"
    rows = ["| Asset | Port | Fingerprint | Action |", "|---|---|---|---|"]
    for item in report.manual_review:
        parts = item.split(None, 2)
        while len(parts) < 3:
            parts.append("–")
        asset, port, rest = parts
        fingerprint, _, action = rest.partition(" — ")
        rows.append(
            "| {} | {} | {} | {} |".format(
                md_escape(asset),
                md_escape(port),
                md_escape(fingerprint),
                md_escape(action) or "–",
            )
        )
    return "\n".join(rows)


def render_markdown(report: Report, title: str | None = None) -> str:
    """Render *report* as a Markdown document."""
    template = Template(TEMPLATE_PATH.read_text(encoding="utf-8"))
    return template.safe_substitute(
        title=title or "Penetration test report — {}".format(report.client),
        client=report.client,
        engagement=report.engagement or "Not specified",
        tester=report.tester or "Not specified",
        date=report.date or datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        generated_at=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        version=__version__,
        exec_summary=_exec_summary(report),
        scope=_scope(report),
        findings=_findings(report),
        host_inventory=_host_inventory(report),
        web_surface=_web_surface(report),
        manual_review=_manual_review(report),
        methodology=METHODOLOGY,
        disclaimer=DISCLAIMER_FULL,
    )


def render_json(report: Report, indent: int = 2) -> str:
    """Render *report* as machine-readable JSON for downstream tooling."""
    payload = report.to_dict()
    payload["generator"] = "reconkit {}".format(__version__)
    payload["generated_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    return json.dumps(payload, indent=indent, ensure_ascii=False)


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #

SEVERITY_COLORS = {
    "critical": "#d73a49",
    "high": "#e36209",
    "medium": "#d4a72c",
    "low": "#1a7f37",
    "info": "#57606a",
}


def html_escape(value: str) -> str:
    """Escape text for safe inclusion in HTML."""
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def _sev_badge(severity: str) -> str:
    severity = severity if severity in SEVERITIES else "info"
    return '<span class="sev sev-{}">{}</span>'.format(severity, severity.upper())


def _html_meta(report: Report) -> str:
    rows = [
        ("Client / owner", report.client),
        ("Engagement", report.engagement or "Not specified"),
        ("Tested by", report.tester or "Not specified"),
        ("Testing date", report.date or datetime.now(timezone.utc).strftime("%Y-%m-%d")),
        ("Report generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Generator", "reconkit {}".format(__version__)),
    ]
    cells = "".join(
        "<tr><td>{}</td><td>{}</td></tr>".format(html_escape(k), html_escape(v)) for k, v in rows
    )
    return '<table class="meta">{}</table>'.format(cells)


def _html_exec_summary(report: Report) -> str:
    counts = report.counts_by_severity()
    total = len(report.findings)

    if total == 0:
        headline = (
            "No findings were raised by the automated checks. This does not mean the environment "
            "is free of vulnerabilities; treat the absence of findings as a prompt for manual "
            "testing."
        )
    else:
        headline = (
            "{total} finding(s) were raised across {hosts} host(s) and {targets} HTTP endpoint(s). "
            "The highest severity observed is {top}. Address critical and high findings first."
        ).format(
            total=total,
            hosts=len(report.hosts),
            targets=len(report.web_targets),
            top=report.findings[0].severity.upper(),
        )

    rows = "".join(
        "<tr><td>{}</td><td>{}</td></tr>".format(severity.capitalize(), counts.get(severity, 0))
        for severity in SEVERITIES
    )
    rows += "<tr><td><b>Total</b></td><td><b>{}</b></td></tr>".format(total)

    bits = [
        "<p><b>Observed attack surface</b></p><ul>",
        "<li>Hosts reported up: <b>{}</b></li>".format(len(report.hosts)),
        "<li>Open ports: <b>{}</b></li>".format(sum(len(h.open_ports) for h in report.hosts)),
        "<li>HTTP endpoints probed: <b>{}</b></li>".format(len(report.web_targets)),
        "</ul>",
    ]
    top = [f for f in report.findings if f.severity in ("critical", "high")][:5]
    if top:
        bits.append("<p><b>Priority items</b></p><ul>")
        for f in top:
            bits.append(
                "<li><code>{}</code> — {} <span class=\"muted\">({})</span></li>".format(
                    html_escape(f.asset), html_escape(f.title), f.severity
                )
            )
        bits.append("</ul>")

    return "<p>{}</p><table><tr><th>Severity</th><th>Count</th></tr>{}</table>{}".format(
        html_escape(headline), rows, "".join(bits)
    )


def _html_scope(report: Report) -> str:
    if not report.scope:
        return (
            '<p class="muted">Scope was not supplied. Populate this section before the report '
            "leaves the team.</p>"
        )
    items = "".join("<li><code>{}</code></li>".format(html_escape(item)) for item in report.scope)
    return "<ul>{}</ul>".format(items)


def _html_finding(index: int, finding: Finding) -> str:
    severity = finding.severity if finding.severity in SEVERITIES else "info"
    parts = [
        '<div class="finding {}">'.format(severity),
        "<h3>{} — {}</h3>".format(index, html_escape(finding.title)),
        "<p>{} <code>{}</code></p>".format(_sev_badge(severity), html_escape(finding.check_id)),
        '<div class="kv"><b>Affected asset:</b> <code>{}</code></div>'.format(
            html_escape(finding.asset)
        ),
    ]
    if finding.description:
        parts.append(
            '<div class="kv"><b>Description:</b> {}</div>'.format(html_escape(finding.description))
        )
    if finding.evidence:
        parts.append(
            '<div class="kv"><b>Evidence:</b></div><div class="evidence">{}</div>'.format(
                html_escape(finding.evidence)
            )
        )
    if finding.remediation:
        parts.append(
            '<div class="kv"><b>Remediation:</b> {}</div>'.format(html_escape(finding.remediation))
        )
    if finding.references:
        parts.append(
            '<div class="kv"><b>References:</b> {}</div>'.format(
                ", ".join(html_escape(ref) for ref in finding.references)
            )
        )
    parts.append("</div>")
    return "".join(parts)


def _html_findings(report: Report) -> str:
    if not report.findings:
        return '<p class="muted">No findings were raised by the automated check set.</p>'
    groups = {severity: [] for severity in SEVERITIES}
    for finding in report.findings:
        groups[finding.severity if finding.severity in SEVERITIES else "info"].append(finding)
    sections = []
    index = 1
    for severity in SEVERITIES:
        if not groups[severity]:
            continue
        sections.append("<h2>{}</h2>".format(severity.capitalize()))
        for finding in groups[severity]:
            sections.append(_html_finding(index, finding))
            index += 1
    return "".join(sections)


def _html_host_inventory(report: Report) -> str:
    if not report.hosts:
        return '<p class="muted">No hosts were parsed from the supplied scan data.</p>'
    rows = ["<tr><th>Host</th><th>Hostname</th><th>Port</th><th>Service</th><th>Banner</th></tr>"]
    for host in report.hosts:
        if not host.open_ports:
            rows.append(
                "<tr><td>{}</td><td>{}</td><td>–</td><td>–</td><td>no open ports</td></tr>".format(
                    html_escape(host.address), html_escape(host.hostname) or "–"
                )
            )
            continue
        for port in host.open_ports:
            rows.append(
                "<tr><td>{}</td><td>{}</td><td>{}/{}</td><td>{}</td><td>{}</td></tr>".format(
                    html_escape(host.address),
                    html_escape(host.hostname) or "–",
                    port.number,
                    port.protocol,
                    html_escape(port.service) or "unknown",
                    html_escape(port.banner),
                )
            )
    return "<table>{}</table>".format("".join(rows))


def _html_web_surface(report: Report) -> str:
    if not report.web_targets:
        return '<p class="muted">No HTTP endpoints were parsed from the supplied httpx data.</p>'
    rows = [
        "<tr><th>URL</th><th>Status</th><th>Title</th><th>Server</th><th>Technologies</th></tr>"
    ]
    for target in report.web_targets:
        rows.append(
            "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                html_escape(target.url),
                target.status_code if target.status_code is not None else "–",
                html_escape(target.title) or "–",
                html_escape(target.webserver) or "–",
                html_escape(", ".join(target.technologies)) or "–",
            )
        )
    return "<table>{}</table>".format("".join(rows))


def _html_manual_review(report: Report) -> str:
    if not report.manual_review:
        return '<p class="muted">No version fingerprints were collected.</p>'
    rows = ["<tr><th>Asset</th><th>Port</th><th>Fingerprint</th><th>Action</th></tr>"]
    for item in report.manual_review:
        parts = item.split(None, 2)
        while len(parts) < 3:
            parts.append("–")
        asset, port, rest = parts
        fingerprint, _, action = rest.partition(" — ")
        rows.append(
            "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                html_escape(asset),
                html_escape(port),
                html_escape(fingerprint),
                html_escape(action) or "–",
            )
        )
    return "<table>{}</table>".format("".join(rows))


_HTML_METHODOLOGY = (
    "<p>Data was collected with standard, non-destructive discovery tooling, normalised and "
    "analysed by <code>reconkit</code>.</p>"
    "<p><b>Collection</b></p><ul>"
    "<li>TCP port and service discovery: <code>nmap -sV -sC</code></li>"
    "<li>HTTP probing and fingerprinting: <code>httpx -jsonl</code></li>"
    "<li>Template-based detection: <code>nuclei -jsonl</code></li>"
    "</ul><p><b>Limitations</b></p><ul>"
    "<li>Version fingerprints are recorded for manual review; reconkit does not claim CVE matches.</li>"
    "<li>No authenticated testing was performed unless stated in the scope.</li>"
    "<li>Automated checks produce false positives and false negatives; validate findings manually.</li>"
    "</ul>"
)


def render_html(report: Report, title: str | None = None) -> str:
    """Render *report* as a self-contained HTML document."""
    doc_title = title or "Penetration test report — {}".format(report.client)
    body = "".join(
        [
            "<h1>{}</h1>".format(html_escape(doc_title)),
            _html_meta(report),
            "<h2>1. Executive summary</h2>",
            _html_exec_summary(report),
            "<h2>2. Scope</h2>",
            _html_scope(report),
            "<h2>3. Findings</h2>",
            _html_findings(report),
            "<h2>4. Asset inventory</h2>",
            _html_host_inventory(report),
            "<h2>5. Web surface</h2>",
            _html_web_surface(report),
            "<h2>6. Manual review required</h2>",
            _html_manual_review(report),
            "<h2>7. Methodology and limitations</h2>",
            _HTML_METHODOLOGY,
            "<h2>8. Disclaimer</h2>",
            '<div class="callout">{}</div>'.format(html_escape(DISCLAIMER_FULL)),
        ]
    )
    template = Template(HTML_TEMPLATE_PATH.read_text(encoding="utf-8"))
    return template.safe_substitute(title=html_escape(doc_title), body=body)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

_SEV_RGB = {
    "critical": (0xD7, 0x3A, 0x49),
    "high": (0xE3, 0x62, 0x09),
    "medium": (0xD4, 0xA7, 0x2C),
    "low": (0x1A, 0x7F, 0x37),
    "info": (0x57, 0x60, 0x6A),
}


def _docx_kv(doc, key: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(key + ": ")
    run.bold = True
    paragraph.add_run(value)


def _docx_table(doc, headers, rows) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)


def _docx_finding(doc, index: int, finding: Finding) -> None:
    from docx.shared import RGBColor

    severity = finding.severity if finding.severity in SEVERITIES else "info"
    heading = doc.add_heading(level=2)
    run = heading.add_run("{} — {}".format(index, finding.title))
    red, green, blue = _SEV_RGB[severity]
    run.font.color.rgb = RGBColor(red, green, blue)
    _docx_kv(doc, "Severity", severity.upper())
    _docx_kv(doc, "Check ID", finding.check_id)
    _docx_kv(doc, "Affected asset", finding.asset)
    if finding.description:
        _docx_kv(doc, "Description", finding.description)
    if finding.evidence:
        _docx_kv(doc, "Evidence", finding.evidence)
    if finding.remediation:
        _docx_kv(doc, "Remediation", finding.remediation)
    if finding.references:
        _docx_kv(doc, "References", ", ".join(finding.references))


def render_docx(report: Report, title: str | None = None) -> bytes:
    """Render *report* as a ``.docx`` byte stream.

    Requires the optional ``python-docx`` dependency: ``pip install 'reconkit[docx]'``.
    """
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise RuntimeError(
            "DOCX export requires python-docx; install it with: pip install 'reconkit[docx]'"
        ) from exc

    doc_title = title or "Penetration test report — {}".format(report.client)
    doc = Document()

    doc.add_heading(doc_title, level=0)

    _docx_table(
        doc,
        ["Field", "Value"],
        [
            ["Client / owner", report.client],
            ["Engagement", report.engagement or "Not specified"],
            ["Tested by", report.tester or "Not specified"],
            ["Testing date", report.date or datetime.now(timezone.utc).strftime("%Y-%m-%d")],
            ["Generator", "reconkit {}".format(__version__)],
        ],
    )

    doc.add_heading("1. Executive summary", level=1)
    counts = report.counts_by_severity()
    total = len(report.findings)
    if total:
        doc.add_paragraph(
            "{} finding(s) across {} host(s) and {} HTTP endpoint(s); highest severity {}.".format(
                total,
                len(report.hosts),
                len(report.web_targets),
                report.findings[0].severity.upper(),
            )
        )
    else:
        doc.add_paragraph("No findings were raised by the automated checks.")
    _docx_table(
        doc,
        ["Severity", "Count"],
        [[severity.capitalize(), str(counts.get(severity, 0))] for severity in SEVERITIES]
        + [["Total", str(total)]],
    )

    doc.add_heading("2. Scope", level=1)
    if report.scope:
        for item in report.scope:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("Scope was not supplied on the command line.")

    doc.add_heading("3. Findings", level=1)
    if report.findings:
        index = 1
        for severity in SEVERITIES:
            group = [
                f
                for f in report.findings
                if (f.severity if f.severity in SEVERITIES else "info") == severity
            ]
            if not group:
                continue
            doc.add_heading(severity.capitalize(), level=2)
            for finding in group:
                _docx_finding(doc, index, finding)
                index += 1
    else:
        doc.add_paragraph("No findings were raised by the automated check set.")

    doc.add_heading("4. Asset inventory", level=1)
    if report.hosts:
        rows = []
        for host in report.hosts:
            for port in host.open_ports:
                rows.append(
                    [
                        host.address,
                        host.hostname or "",
                        "{}/{}".format(port.number, port.protocol),
                        port.service or "unknown",
                        port.banner,
                    ]
                )
        _docx_table(doc, ["Host", "Hostname", "Port", "Service", "Banner"], rows)
    else:
        doc.add_paragraph("No hosts were parsed.")

    doc.add_heading("5. Web surface", level=1)
    if report.web_targets:
        rows = [
            [
                target.url,
                str(target.status_code) if target.status_code is not None else "",
                target.title,
                target.webserver,
                ", ".join(target.technologies),
            ]
            for target in report.web_targets
        ]
        _docx_table(doc, ["URL", "Status", "Title", "Server", "Technologies"], rows)
    else:
        doc.add_paragraph("No HTTP endpoints were parsed.")

    doc.add_heading("6. Manual review required", level=1)
    if report.manual_review:
        for item in report.manual_review:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("No version fingerprints were collected.")

    doc.add_heading("7. Methodology and limitations", level=1)
    doc.add_paragraph(
        "Data was collected with standard, non-destructive discovery tooling, normalised and "
        "analysed by reconkit. Version fingerprints are recorded for manual review only; "
        "reconkit does not claim CVE matches, and automated checks produce both false positives "
        "and false negatives."
    )

    doc.add_heading("8. Disclaimer", level=1)
    for paragraph in DISCLAIMER_FULL.strip().split("\n\n"):
        doc.add_paragraph(paragraph)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
